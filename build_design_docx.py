# -*- coding: utf-8 -*-
"""Build 一戶通智能日曆 設計摘要 DOCX (standard_business_brief preset)."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

EAST = "Microsoft JhengHei"
LATIN = "Calibri"
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x0B, 0x25, 0x45)
GRAY = RGBColor(0x59, 0x59, 0x59)
LINK = RGBColor(0x05, 0x63, 0xC1)


def set_run_font(run, latin=LATIN, east=EAST, size=None, bold=None, color=None):
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def patch_style(style, size, color, before, after, line=None, bold=True):
    style.font.name = LATIN
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), LATIN)
    rFonts.set(qn("w:hAnsi"), LATIN)
    rFonts.set(qn("w:eastAsia"), EAST)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line


def add_numbering(doc):
    """Append bullet/decimal definitions to the existing numbering part."""
    numbering = doc.part.numbering_part._element
    nsmap = "xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'"

    def abstract(num_id, fmt, text):
        return (
            f'<w:abstractNum w:abstractNumId="{num_id}" {nsmap}>'
            '<w:multiLevelType w:val="hybridMultilevel"/>'
            '<w:lvl w:ilvl="0">'
            '<w:start w:val="1"/>'
            f'<w:numFmt w:val="{fmt}"/>'
            f'<w:lvlText w:val="{text}"/>'
            '<w:lvlJc w:val="left"/>'
            "<w:pPr><w:tabs><w:tab w:val=\"num\" w:pos=\"720\"/></w:tabs>"
            '<w:ind w:left="720" w:hanging="360"/></w:pPr>'
            "</w:lvl></w:abstractNum>"
        )

    from docx.oxml import parse_xml
    for xml in (
        abstract(100, "bullet", "\u2022"),
        abstract(101, "decimal", "%1."),
        f'<w:num w:numId="100" {nsmap}><w:abstractNumId w:val="100"/></w:num>',
        f'<w:num w:numId="101" {nsmap}><w:abstractNumId w:val="101"/></w:num>',
    ):
        numbering.append(parse_xml(xml))


def add_list_para(doc, text, num_id=100, bold_prefix=None, size=11):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.167
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), str(num_id))
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        set_run_font(r, size=size, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def add_hyperlink(paragraph, url, text, size=9):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), LATIN)
    rFonts.set(qn("w:hAnsi"), LATIN)
    rFonts.set(qn("w:eastAsia"), EAST)
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size * 2))
    rPr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def make_table(doc, headers, rows, widths, header_fill="F2F4F7"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "9360")
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    tblInd = OxmlElement("w:tblInd")
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    tblPr.append(tblInd)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "BFBFBF")
        borders.append(e)
    tblPr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for m, v in (("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")):
        e = OxmlElement("w:" + m)
        e.set(qn("w:w"), v)
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)
    for gc, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Emu(w * 635)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    hdr = table.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    for i, cell in enumerate(hdr.cells):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), header_fill)
        tcPr.append(shd)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(headers[i])
        set_run_font(r, size=10, bold=True)
    for ridx, row_data in enumerate(rows, start=1):
        for cidx, text in enumerate(row_data):
            cell = table.rows[ridx].cells[cidx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(text)
            set_run_font(r, size=10)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=9, color=GRAY)
    return p


def add_body(doc, text, size=11, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def add_h(doc, text, level):
    h = doc.add_heading(text, level=level)
    return h


def build():
    doc = Document()
    # Base styles
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(11)
    rPr = normal.element.get_or_add_rPr()
    rPr.get_or_add_rFonts().set(qn("w:eastAsia"), EAST)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    patch_style(doc.styles["Heading 1"], 16, BLUE, 16, 8)
    patch_style(doc.styles["Heading 2"], 13, BLUE, 12, 6)
    patch_style(doc.styles["Heading 3"], 12, DARK_BLUE, 8, 4)
    add_numbering(doc)

    # Footer: label + page number
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run("一戶通智能日曆 — 項目設計摘要")
    set_run_font(r, size=8.5, color=GRAY)
    r = fp.add_run("\t")
    set_run_font(r, size=8.5, color=GRAY)
    run = fp.add_run("第 ")
    set_run_font(run, size=8.5, color=GRAY)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r = fp.add_run()
    set_run_font(r, size=8.5, color=GRAY)
    r._element.append(fld1)
    r = fp.add_run()
    set_run_font(r, size=8.5, color=GRAY)
    r._element.append(instr)
    r = fp.add_run()
    set_run_font(r, size=8.5, color=GRAY)
    r._element.append(fld2)
    r = fp.add_run(" 頁")
    set_run_font(r, size=8.5, color=GRAY)

    # Title block
    tp = doc.add_paragraph()
    tp.paragraph_format.space_after = Pt(2)
    r = tp.add_run("一戶通智能日曆 — 項目設計摘要")
    set_run_font(r, size=22, bold=True, color=INK)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    r = sp.add_run("澳門創新比賽項目設計文件")
    set_run_font(r, size=12, color=GRAY)
    mp = doc.add_paragraph()
    mp.paragraph_format.space_after = Pt(8)
    r = mp.add_run("版本 1.0　|　日期：2026-08-01　|　團隊：3 人　|　形態：Web　|　介面語言：繁體中文")
    set_run_font(r, size=9.5, color=GRAY)
    pPr = mp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "2E74B5")
    pBdr.append(bottom)
    pPr.append(pBdr)

    # 1. 項目概述
    add_h(doc, "一、項目概述", 1)
    add_body(
        doc,
        "一戶通智能日曆是一款面向澳門居民的 AI 日曆。用戶上傳政府通知、證件或任何含日期的文檔，"
        "AI 自動抽出「日期＋事件」並寫入日曆；用戶亦可用自然語言搜尋任意主題（例如「我想要英超賽程」），"
        "AI 將含日期錨點的結果填入日曆。所有事件均可回溯來源：浮窗顯示原文高光段落，輕擊查看完整原檔或網站。",
    )
    add_list_para(doc, "一般澳門居民", bold_prefix="目標用戶：")
    add_list_para(doc, "AI 讀文檔 → 日曆；AI 即時搜尋 → 日曆；到期提醒", bold_prefix="核心能力：")
    add_list_para(doc, "文檔 → 日曆（主秀）；AI 搜尋與到期提醒（輔）", bold_prefix="Demo 主秀：")

    # 2. API 研究結果
    add_h(doc, "二、一戶通 API 研究結果", 1)
    add_body(
        doc,
        "已查證 ga.gov.mo 公開文檔（2026-08-01）。結論：官方存在「一戶通自然人帳戶 OAuth2 API」，"
        "可取得證件到期日；津貼到期日則無任何公開 API。",
    )
    make_table(
        doc,
        ["端點", "用途", "說明"],
        [
            ["GET /o/authorize/", "授權請求", "OAuth 2.0 Authorization Code + PKCE"],
            ["POST /o/token/", "交換訪問碼", "以授權碼換取 access_token"],
            ["GET /o/profile/", "用戶資料", "scope=profile；返回 euid、mobile、identityDocs"],
        ],
        [2050, 2000, 5310],
    )
    add_caption(doc, "主機：account.gov.mo（規格 v1.7.2，2024-01）。")
    add_list_para(doc, "identityDocs 含 identityValidDate（證件有效日期），即「身份證到期日」可透過官方 API 取得")
    add_list_para(doc, "需向行政公職局（電子政務廳）申請 client_id／client_secret；申請表樣本對象為政府部門，學生隊伍可否申請需直接查詢（8866 8866，需驗證）")
    add_list_para(doc, "其他模組接口：/identity、/qrcode-check、/file-processing、/progress、/notification、/my-photo、/app/v1.0/memo；詳細文檔位於 G2E 內部知識庫，公眾不可見")
    add_list_para(doc, "津貼到期日：無公開 API（需驗證部門級接口），故採手動輸入＋文檔抽取")
    add_list_para(doc, "澳門《個人資料保護法》適用，處理個人資料需用戶明確同意")

    # 3. 已定決策
    add_h(doc, "三、已定決策", 1)
    make_table(
        doc,
        ["#", "決策", "出處"],
        [
            ["1", "政府資料混合路線：證件到期日走一戶通 OAuth API；津貼到期日手動輸入＋OCR；demo 用模擬數據", "ADR 0001"],
            ["2", "Web 優先，架構預留遷移至手機 app", "ADR 0002"],
            ["3", "自建日曆，不同步 Google／Apple 日曆", "ADR 0003"],
            ["4", "事件存 localStorage、原檔存 IndexedDB、金鑰與 AI 處理放薄後端", "ADR 0004"],
            ["5", "文檔事件全自動寫入：來源可追溯＋一鍵復原，無確認步驟", "ADR 0005"],
        ],
        [620, 6360, 2380],
    )
    add_caption(doc, "完整決策文件存放於本項目 docs/adr/ 目錄。")

    # 4. 產品規格
    add_h(doc, "四、產品規格", 1)
    add_h(doc, "4.1 文檔事件抽取", 2)
    add_list_para(doc, "Word、PDF、Excel、圖片／截圖", bold_prefix="支援格式：")
    add_list_para(doc, "視為批量事件表：日期欄＋標題欄，一行一個事件", bold_prefix="Excel：")
    add_list_para(doc, "抽出「日期＋事件描述」配對；OCR 支援繁體中文、葡文、英文", bold_prefix="Word／PDF／圖片：")
    add_list_para(doc, "無日期錨點的內容不構成日曆事件", bold_prefix="規則：")
    add_h(doc, "4.2 AI 即時搜尋", 2)
    add_list_para(doc, "任意主題；有明確日期的結果自動寫入日曆並帶來源 URL", bold_prefix="行為：")
    add_list_para(doc, "回覆「找不到可排程內容」，不入日曆", bold_prefix="無日期結果：")
    add_h(doc, "4.3 到期提醒", 2)
    add_list_para(doc, "瀏覽器通知", bold_prefix="渠道：")
    add_list_para(doc, "提前 90／30／7 天提醒；津貼提前 7 天；一般事件前一天（可自訂）", bold_prefix="節奏：")
    add_h(doc, "4.4 互動", 2)
    add_list_para(doc, "浮窗顯示來源原文高光段落", bold_prefix="電腦 hover／手機長按：")
    add_list_para(doc, "app 內檢視完整原檔（Word／PDF／Excel／圖片）或開啟網站", bold_prefix="輕擊：")
    add_list_para(doc, "每次匯入為一個批次，可一鍵撤銷", bold_prefix="防護：")
    add_h(doc, "4.5 Demo 主秀", 2)
    add_list_para(doc, "文檔 → 日曆；配角為 AI 搜尋與到期提醒；一戶通登入用模擬數據", bold_prefix="主秀：")

    # 5. 術語表
    add_h(doc, "五、術語表", 1)
    make_table(
        doc,
        ["術語", "定義"],
        [
            ["一戶通", "澳門特區政府統一電子帳戶（Macao One Account），用戶可藉其登入政府電子服務。避免誤用「澳門通」。"],
            ["用戶", "使用本應用的自然人（澳門居民或外地僱員）。"],
            ["證件到期日", "身份證明文件的到期日期，透過一戶通 OAuth API 的 identityDocs.identityValidDate 取得。"],
            ["津貼到期日", "政府津貼或優惠的到期日期；目前無官方 API，由用戶手動輸入或文檔抽取。"],
            ["文檔事件抽取", "AI 從文檔抽出「日期＋事件描述」配對並寫入日曆的流程。"],
            ["事件來源", "產生事件的文檔與原文引用；浮窗高光顯示，輕擊可看完整原檔或網站。"],
            ["匯入批次", "一次文檔匯入產生的整組事件，可整體撤銷。"],
            ["日曆事件", "有明確日期錨點的內容單元；無日期內容不構成日曆事件。"],
        ],
        [2050, 7310],
    )

    # 6. 待定項目
    add_h(doc, "六、待定項目", 1)
    add_list_para(doc, "比賽評分標準（決定 demo 主秀最終配比）")
    add_list_para(doc, "技術棧（建議組合：Next.js ＋ Vercel Functions ＋ OpenAI API ＋ pdf.js／mammoth／SheetJS）")
    add_list_para(doc, "產品名稱")
    add_list_para(doc, "一戶通 client_id 申請資格（聯絡行政公職局 8866 8866）")

    # 7. 風險與對策
    add_h(doc, "七、風險與對策", 1)
    make_table(
        doc,
        ["風險", "對策"],
        [
            ["AI 抽錯日期", "全自動寫入搭配來源可追溯＋一鍵復原"],
            ["一戶通 3.0 引入 AI 主動服務", "差異化：第三方、跨資料來源、用戶自主控制"],
            ["個人資料保護法", "demo 全程模擬數據；正式上線需用戶明確同意"],
            ["中葡雙語 OCR 品質", "多模態 LLM 直接讀原檔，繞過傳統 OCR 語言瓶頸"],
            ["原檔儲存容量", "用 IndexedDB（localStorage 上限約 5MB）"],
        ],
        [2800, 6560],
    )

    # 8. 附錄
    add_h(doc, "八、附錄：官方資料", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("1. ")
    set_run_font(r, size=10)
    add_hyperlink(
        p,
        "https://ga.gov.mo/macao-ga-portal/_spaces/GTDSEPPG/2024/%e4%b8%80%e6%88%b6%e9%80%9a%e8%87%aa%e7%84%b6%e4%ba%ba%e5%b8%b3%e6%88%b6%e6%8a%80%e8%a1%93%e8%a6%8f%e6%a0%bc%e8%a6%81%e6%b1%82v1.7.2.pdf",
        "一戶通自然人帳戶技術規格要求 v1.7.2（行政公職局）",
        size=10,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("2. ")
    set_run_font(r, size=10)
    add_hyperlink(
        p,
        "https://ga.gov.mo/macao-ga-portal/_spaces/GTDSEPPG/2024/%e4%b8%80%e6%88%b6%e9%80%9a2.0%e6%9c%8d%e5%8b%99%e6%95%b4%e5%90%88%e7%94%b3%e8%ab%8b%e8%a1%a8v202309%e5%8f%83%e8%80%83%e6%a8%a3%e6%9c%ac.pdf",
        "一戶通 2.0 服務整合申請表 v202309 參考樣本",
        size=10,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("3. ")
    set_run_font(r, size=10)
    add_hyperlink(
        p,
        "https://www.gov.mo/zh-hans/services/ps-2112/",
        "澳門政府服務網 — 一戶通（實體）",
        size=10,
    )

    out = "一戶通智能日曆_設計摘要.docx"
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    build()
